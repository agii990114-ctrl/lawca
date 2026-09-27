"""서식 DOCX 템플릿(docxtpl)을 만든다. 결과는 src/lawca/forms/data/*.docx로 저장해 커밋한다.

법인이 쓰던 양식으로 바꾸고 싶으면 같은 자리표시자({{ case_number }} 등)를 넣은 DOCX로 파일을 바꾸면 된다.

사용법: uv run python scripts/make_form_templates.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).resolve().parents[1] / "src" / "lawca" / "forms" / "data"
FONT = "바탕"


def new_document(title: str) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.core_properties.title = title
    doc.core_properties.author = "lawca"
    doc.core_properties.comments = "lawca가 만든 초안입니다. 제출 전에 담당 변호사가 검토해야 합니다."
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(" ".join(title))
    run.bold = True
    run.font.size = Pt(20)
    doc.add_paragraph()
    return doc


def line(doc: Document, text: str, *, align: WD_ALIGN_PARAGRAPH | None = None, space_after: int = 6) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        paragraph.alignment = align


def case_block(doc: Document) -> None:
    line(doc, "사        건    {{ case_number }}  {{ case_name }}")
    line(doc, "원        고    {{ plaintiffs }}")
    line(doc, "피        고    {{ defendants }}")
    doc.add_paragraph()


def signature(doc: Document) -> None:
    doc.add_paragraph()
    line(doc, "{{ filed_on }}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    line(doc, "신청인  {{ applicant }}  (서명 또는 날인)", align=WD_ALIGN_PARAGRAPH.RIGHT)
    line(doc, "{%p if agent %}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    line(doc, "소송대리인  {{ agent }}  (서명 또는 날인)", align=WD_ALIGN_PARAGRAPH.RIGHT)
    line(doc, "{%p endif %}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_paragraph()
    heading = doc.add_paragraph()
    run = heading.add_run("{{ court }}  귀중")
    run.bold = True
    run.font.size = Pt(14)


def certificate_of_finality() -> Document:
    doc = new_document("확정증명원")
    case_block(doc)
    line(doc, "위 사건에 관하여 {{ judgment_date }} 선고한 판결은 {{ finality_date }} 확정되었음을 증명하여 주시기 바랍니다.")
    signature(doc)
    return doc


def certificate_of_service() -> Document:
    doc = new_document("송달증명원")
    case_block(doc)
    line(doc, "위 사건에 관하여 {{ served_document }}이(가) {{ served_party }}에게 송달되었음을 증명하여 주시기 바랍니다.")
    line(doc, "신청 통수  {{ copies }}통")
    signature(doc)
    return doc


def address_correction() -> Document:
    doc = new_document("주소보정서")
    case_block(doc)
    line(doc, "위 사건에 관하여 {{ target_party }}에 대한 소송서류가 송달되지 않았으므로 다음과 같이 보정합니다.")
    doc.add_paragraph()
    line(doc, "신청 내용  {{ method }}")
    line(doc, "{%p if new_address %}")
    line(doc, "새 주소  {{ new_address }}")
    line(doc, "{%p endif %}")
    line(doc, "{%p if attachments %}")
    line(doc, "첨부 서류  {{ attachments }}")
    line(doc, "{%p endif %}")
    signature(doc)
    return doc


def fact_inquiry() -> Document:
    doc = new_document("사실조회신청서")
    case_block(doc)
    line(doc, "위 사건에 관하여 신청인은 주장사실을 증명하기 위하여 다음과 같이 사실조회를 신청합니다.")
    doc.add_paragraph()
    line(doc, "1. 사실조회 기관의 명칭과 주소")
    line(doc, "    명칭  {{ institution }}")
    line(doc, "    주소  {{ institution_address }}")
    line(doc, "2. 증명하려는 사실")
    line(doc, "    {{ purpose }}")
    line(doc, "3. 사실조회 사항")
    line(doc, "    {{ inquiry_items }}")
    signature(doc)
    return doc


def execution_clause() -> Document:
    doc = new_document("집행문부여신청")
    case_block(doc)
    line(doc, "채  권  자    {{ creditor }}")
    line(doc, "채  무  자    {{ debtor }}")
    doc.add_paragraph()
    line(doc, "위 사건에 관하여 {{ judgment_date }}자 {{ title_document }} 정본에 집행문을 부여하여 주시기 바랍니다.")
    line(doc, "신청 통수  {{ copies }}통")
    signature(doc)
    return doc


def brief() -> Document:
    """준비서면 틀. 본문은 변호사가 쓴 글(문단 목록)이고, 입증방법·첨부서류는 증거 목록에서 채운다."""
    doc = new_document("준비서면")
    heading = doc.paragraphs[0]
    heading.runs[0].text = "{{ title_spaced }}"  # "원고 제2준비서면"처럼 제목을 바꿀 수 있다
    case_block(doc)
    line(doc, "위 사건에 관하여 {{ subject }} 다음과 같이 변론을 준비합니다.")
    doc.add_paragraph()
    line(doc, "다        음", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "{%p for para in body %}")
    line(doc, "{{ para }}", space_after=10)
    line(doc, "{%p endfor %}")
    doc.add_paragraph()
    line(doc, "{%p if evidence %}")
    line(doc, "입  증  방  법", align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "{%p for e in evidence %}")
    line(doc, "1. {{ e.label }}    {{ e.title }}")
    line(doc, "{%p endfor %}")
    line(doc, "{%p endif %}")
    line(doc, "{%p if attachments %}")
    line(doc, "첨  부  서  류", align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "{%p for a in attachments %}")
    line(doc, "1. {{ a }}")
    line(doc, "{%p endfor %}")
    line(doc, "{%p endif %}")
    doc.add_paragraph()
    line(doc, "{{ filed_on }}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    line(doc, "{{ side }} {{ representative }}  {{ signer }}  (서명 또는 날인)", align=WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_paragraph()
    court = doc.add_paragraph()
    run = court.add_run("{{ court }}  귀중")
    run.bold = True
    run.font.size = Pt(14)
    return doc


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, build in [
        ("certificate_of_finality", certificate_of_finality),
        ("certificate_of_service", certificate_of_service),
        ("address_correction", address_correction),
        ("fact_inquiry", fact_inquiry),
        ("execution_clause", execution_clause),
        ("brief", brief),
    ]:
        path = OUT / f"{name}.docx"
        build().save(path)
        print(path)


if __name__ == "__main__":
    main()
