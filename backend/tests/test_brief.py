"""상대방 서면 요약, 증거 목록, 준비서면 틀, 사건 API 테스트."""

import io
from pathlib import Path

import pytest
from docx import Document

from lawca.agent.brief import BriefSummary, CitedEvidence, Claim, evidence_label, number_key, parse_label, summary_out
from lawca.api.app import app, get_extractor_factory
from lawca.extraction.schema import DocumentType, Party, TextField
from tests.fakes import FakeChatModel
from tests.test_agent import use_models
from tests.test_api import FakeExtractor, chat, client, new_conversation  # noqa: F401 (픽스처 재사용)
from tests.test_extraction import correction_order, ev

ANSWER_PDF = Path(__file__).parent / "fixtures" / "synthetic_answer.pdf"
PARTIES = [Party(role="원고", name="홍길동", evidence=ev("원 고 홍길동")), Party(role="피고", name="김철수", evidence=ev("피 고 김철수"))]

SUMMARY = BriefSummary(
    submitter="피고",
    request_summary="원고의 청구를 기각한다는 판결을 구함",
    claims=[
        Claim(point="돈을 받은 사실은 인정", detail="3,000만 원 수령 인정", quote="피고가 원고로부터 3,000만 원을 받은 사실은 인정합니다", page=1),
        Claim(point="대여가 아니라 증여", detail="사업을 돕기 위한 증여", quote="빌려준 것이 아닙니다", page=1),
        Claim(point="소멸시효 완성", detail="청구권 소멸", quote="이미 소멸시효가 완성되었습니다", page=1),  # 실제로는 2쪽
        Claim(point="지어낸 주장", detail="원문에 없음", quote="원고는 사기꾼입니다", page=2),
    ],
    evidence=[
        CitedEvidence(label="을 제1호증", title="카카오톡 대화 내역"),
        CitedEvidence(label="을제2호증의1", title="사업자등록증명"),
        CitedEvidence(label="을 제1호증", title="중복"),
        CitedEvidence(label="증인 박가상", title="증인"),
    ],
)


def docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


# 순수 함수


def test_evidence_labels():
    assert parse_label("을 제2호증의1") == ("을", "2-1")
    assert parse_label("갑제10호증") == ("갑", "10")
    assert parse_label("증인 홍길동") is None
    assert evidence_label("을", "2-1") == "을 제2호증의1"
    assert sorted(["10", "2-1", "2", "1"], key=number_key) == ["1", "2", "2-1", "10"]


def test_summary_out_verifies_quotes_and_normalizes_evidence():
    pages = ["피고가 원고로부터 3,000만 원을 받은 사실은 인정합니다. 빌려준 것이 아닙니다.", "청구권은 이미 소멸시효가\n완성되었습니다."]
    out = summary_out(SUMMARY, pages)
    assert [(c["verified"], c["page"]) for c in out["claims"]] == [(True, 1), (True, 1), (True, 2), (False, 2)]
    assert [(e["label"], e["title"]) for e in out["evidence"]] == [("을 제1호증", "카카오톡 대화 내역"), ("을 제2호증의1", "사업자등록증명")]


# 채팅: 상대방 답변서 요약


@pytest.fixture
def answer_chat(client):  # noqa: F811
    doc = correction_order(
        document_type=DocumentType.ANSWER, document_type_evidence=ev("답 변 서"), parties=PARTIES, designated_period=None,
        issued_date=TextField(value="2026-09-20", evidence=ev("2026. 9. 20.")),
    )
    app.dependency_overrides[get_extractor_factory] = lambda: (lambda: FakeExtractor(doc))
    use_models(FakeChatModel(brief=SUMMARY))
    file_id = client.post("/api/files", files={"file": ("답변서.pdf", ANSWER_PDF.read_bytes(), "application/pdf")}).json()["id"]
    return chat(client, new_conversation(client), file_ids=[file_id])


def test_answer_is_summarized_in_chat(client, answer_chat):  # noqa: F811
    card = next(e["card"] for e in answer_chat if e["type"] == "card" and e["card"]["kind"] == "brief_summary")
    assert card["document_type"] == "답변서" and card["case_number"] == "2026가단51234"
    assert [c["verified"] for c in card["claims"]] == [True, True, True, False]
    assert card["claims"][2]["page"] == 2  # 모델이 1쪽이라고 했지만 원문은 2쪽
    text = "".join(e["delta"] for e in answer_chat if e["type"] == "text")
    assert "상대방 주장 4건, 증거 2건" in text and "찾지 못한 주장이 1건" in text
    # 요약은 문서에 저장되고, 답변서에는 기한 후보가 없어 대기 목록에 오르지 않는다
    detail = client.get("/api/cases/2026가단51234").json()
    assert detail["document_list"][0]["summary"]["submitter"] == "피고"
    assert client.get("/api/deadlines/pending").json()["documents"] == []


def test_scanned_answer_is_not_summarized(client):  # noqa: F811
    doc = correction_order(document_type=DocumentType.ANSWER, document_type_evidence=ev("답 변 서"), designated_period=None)
    app.dependency_overrides[get_extractor_factory] = lambda: (lambda: FakeExtractor(doc))
    use_models(FakeChatModel(brief=SUMMARY))
    file_id = client.post("/api/files", files={"file": ("scan.pdf", b"%PDF-1.4 fake", "application/pdf")}).json()["id"]
    events = chat(client, new_conversation(client), file_ids=[file_id])
    assert not any(e["type"] == "card" and e["card"]["kind"] == "brief_summary" for e in events)
    assert "스캔본이라" in "".join(e["delta"] for e in events if e["type"] == "text")


# 증거 목록


def test_evidence_from_summary_and_crud(client, answer_chat):  # noqa: F811
    card = next(e["card"] for e in answer_chat if e["type"] == "card" and e["card"]["kind"] == "brief_summary")
    url = "/api/cases/2026가단51234/evidence"
    body = {"source_document_id": card["document_id"], "items": card["evidence"]}
    assert client.post(f"{url}/bulk", json=body).json() == {"added": ["을 제1호증", "을 제2호증의1"], "skipped": []}
    assert client.post(f"{url}/bulk", json=body).json()["skipped"] == ["을 제1호증", "을 제2호증의1"]

    first = client.post(url, json={"side": "갑", "title": "차용증 사본"}).json()
    second = client.post(url, json={"side": "갑", "title": "계좌이체 내역", "number": "2의1"}).json()
    assert (first["label"], second["label"]) == ("갑 제1호증", "갑 제2호증의1")
    assert client.post(url, json={"side": "갑", "title": "x", "number": "1"}).status_code == 409
    assert client.post(url, json={"side": "갑", "title": "x", "number": "하나"}).status_code == 422
    assert client.post(url, json={"side": "갑", "title": "문자"}).json()["label"] == "갑 제3호증"

    patched = client.patch(f"/api/evidence/{first['id']}", json={"submitted_on": "2026-09-28", "note": "원본 보관"}).json()
    assert patched["submitted_on"] == "2026-09-28" and patched["note"] == "원본 보관"
    assert client.patch(f"/api/evidence/{first['id']}", json={"clear_submitted": True}).json()["submitted_on"] is None
    assert client.delete(f"/api/evidence/{second['id']}").status_code == 204

    detail = client.get("/api/cases/2026가단51234").json()
    assert [e["label"] for e in detail["evidence_list"]] == ["갑 제1호증", "갑 제3호증", "을 제1호증", "을 제2호증의1"]
    assert detail["evidence_list"][2]["from_summary"] is True
    assert client.get("/api/cases", params={"q": "홍길동"}).json()[0]["evidence"] == 4


# 준비서면 본문 초안(변호사)

from lawca.agent.brief_writer import BriefDraft, BriefInputs, Paragraph, Section, body_text, check, note_lines  # noqa: E402
from tests.conftest import login, make_user  # noqa: E402

NOTES = """- 증여 아님: 통장 표시 '대여금'(갑1)
- 문자는 변제 약속
- 소멸시효 중단"""

DRAFT = BriefDraft(
    sections=[
        Section(heading="1. 피고 주장의 요지", paragraphs=[Paragraph(text="피고는 증여라고 주장합니다.", sources=["상대방 서면"])]),
        Section(
            heading="2. 증여 주장에 대한 반박",
            paragraphs=[
                Paragraph(text="원고는 이체할 때 통장 표시를 '대여금'으로 적었습니다(갑 제1호증).", sources=["메모1", "갑 제1호증"]),
                Paragraph(text="피고는 '갚을게'라고 보냈습니다. [인용 확인 필요]", sources=["메모2"]),
            ],
        ),
        Section(heading="3. 결론", paragraphs=[Paragraph(text="원고의 청구는 인용되어야 합니다.", sources=["메모2"])]),
    ],
    open_points=["소멸시효 법리 보충"],
)


def test_note_lines_and_checks():
    assert note_lines(NOTES) == ["증여 아님: 통장 표시 '대여금'(갑1)", "문자는 변제 약속", "소멸시효 중단"]
    inputs = BriefInputs(record="사건 2026가단51234", opponent=["증여"], evidence=["갑 제1호증 계좌이체 내역"], notes=NOTES)
    result = check(inputs, DRAFT)
    assert result["notes_tracked"] is True and [n["number"] for n in result["unused_notes"]] == [3]
    assert result["placeholders"] == 1 and result["case_law"] == [] and result["unknown_evidence"] == []
    bad = DRAFT.model_copy(update={"sections": [Section(heading="x", paragraphs=[Paragraph(
        text="대법원 2020. 1. 9. 선고 2019다12345 판결, 민법 제168조, 갑 제9호증, 500만 원", sources=["메모"])])]})
    result = check(inputs, bad)
    assert result["notes_tracked"] is False and result["unused_notes"] == []  # 번호가 없으면 판단하지 않는다
    assert result["case_law"] and result["statutes_not_in_notes"] == ["민법 제168조"]
    assert result["unknown_evidence"] == ["갑 제9호증"] and result["amounts_not_in_inputs"] == ["500만 원"]
    assert body_text(DRAFT).startswith("1. 피고 주장의 요지\n피고는 증여라고 주장합니다.\n\n2. 증여 주장에 대한 반박\n")
